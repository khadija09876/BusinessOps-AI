import io
import re
import json
import time
import html
from datetime import datetime, timedelta

import streamlit as st
import pandas as pd
import plotly.express as px
from groq import Groq

# CrewAI Flow
from crewai.flow import Flow, start, listen

# Report generation libraries
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from docx import Document


# ============================================================
# BUSINESSOPS AI
# Autonomous Business Process Intelligence Platform
# 6 dynamic agents -> each agent does real work (1 LLM call each)
# ============================================================

st.set_page_config(
    page_title="BusinessOps AI",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="expanded",
)


# ============================================================
# CONFIG
# ============================================================

MODEL = "openai/gpt-oss-20b"

# Streamlit version compatibility (use_container_width -> width="stretch")
try:
    _VER = tuple(int(x) for x in st.__version__.split(".")[:2])
except Exception:
    _VER = (1, 0)
FULL = {"width": "stretch"} if _VER >= (1, 50) else {"use_container_width": True}

HORIZON_HOURS = {
    "Immediate (24-48 Hours)": 48,
    "30 Days (Short-term)": 30 * 24,
    "90 Days (Quarterly)": 90 * 24,
    "6 Months (Strategic)": 180 * 24,
}

HORIZON_SHORT = {
    "Immediate (24-48 Hours)": "24-48 Hrs",
    "30 Days (Short-term)": "30 Days",
    "90 Days (Quarterly)": "90 Days",
    "6 Months (Strategic)": "6 Months",
}

HORIZON_NOTES = {
    "Immediate (24-48 Hours)": "Everything must fit inside 48 hours. Use hour-based milestones, no long-term programs.",
    "30 Days (Short-term)": "Plan across about 4 weeks with weekly milestones and quick wins.",
    "90 Days (Quarterly)": "Plan across one quarter with monthly milestones and a mid-point review.",
    "6 Months (Strategic)": "Plan strategically across 6 months with phased roll-out, governance and quarterly reviews.",
}

PRIORITY_NOTES = {
    "Critical / Urgent": "Treat as an emergency: front-load work, short feedback loops, explicit escalation paths, deadlines in hours, parallel workstreams.",
    "High": "Fast and focused execution: tight milestones, frequent checkpoints, strict scope control.",
    "Medium": "Balanced pace: standard governance, planned checkpoints, moderate resource commitment.",
    "Low": "Low urgency: sequence around other work, minimal resourcing, light governance, cost efficiency first.",
}

LIKELIHOOD = {"low": 1, "medium": 2, "high": 3}
IMPACT = {"low": 1, "medium": 2, "high": 3, "critical": 4}


# ============================================================
# PREMIUM CHARCOAL + DEEP MAROON THEME & UI STYLES
# ============================================================

st.markdown(
    """
    <style>

    /* ======================================================
        GLOBAL APPLICATION
        ====================================================== */

    .stApp {
        background:
            radial-gradient(
                circle at 10% 0%,
                rgba(122, 34, 55, 0.10),
                transparent 28%
            ),
            radial-gradient(
                circle at 90% 5%,
                rgba(255, 255, 255, 0.025),
                transparent 25%
            ),
            linear-gradient(
                135deg,
                #07080a 0%,
                #0b0c0f 48%,
                #090a0d 100%
            );
        color: #eeeeec;
    }

    .main .block-container {
        max-width: 1380px;
        padding-top: 1.2rem;
        padding-bottom: 2.5rem;
    }

    #MainMenu {
        visibility: hidden;
    }

    footer {
        visibility: hidden;
    }

    header,
    [data-testid="stHeader"] {
        background: transparent !important;
    }


    /* ======================================================
        CUSTOM COMPONENTS (BADGE, HERO, STAGES, REPORT)
        ====================================================== */

    .badge {
        display: inline-block;
        padding: 4px 10px;
        background: rgba(122, 34, 55, 0.25);
        border: 1px solid rgba(177, 76, 101, 0.4);
        border-radius: 20px;
        color: #d98294;
        font-size: 10px;
        font-weight: 700;
        letter-spacing: 1px;
        margin-bottom: 12px;
    }

    .hero {
        padding: 10px 0 20px 0;
    }

    .stage {
        background: linear-gradient(145deg, rgba(20, 19, 23, 0.98), rgba(11, 12, 15, 0.98));
        border: 1px solid rgba(255, 255, 255, 0.065);
        border-radius: 14px;
        padding: 20px;
        margin-bottom: 15px;
        box-shadow: 0 10px 25px rgba(0, 0, 0, 0.15);
        height: 100%;
    }

    .stage-number {
        font-size: 11px;
        color: #8c334d;
        font-weight: 750;
        letter-spacing: 0.8px;
        margin-bottom: 6px;
    }

    .stage-title {
        font-size: 15px;
        color: #f1f1ef;
        font-weight: 750;
        margin-bottom: 8px;
    }

    .stage-text {
        font-size: 12px;
        color: #969aa3;
        line-height: 1.6;
    }

    .stage-status {
        font-size: 11px;
        font-weight: 700;
        margin-top: 10px;
        letter-spacing: 0.3px;
    }
    .stage-status.idle { color: #5d6270; }
    .stage-status.ok { color: #6fbf8a; }
    .stage-status.warn { color: #d9a45f; }
    .stage-status.fail { color: #d9626f; }

    .report {
        background: linear-gradient(145deg, #0e1014, #0b0c10);
        border: 1px solid #282b32;
        border-radius: 14px;
        padding: 25px;
        color: #d7d7d4;
        line-height: 1.7;
        font-size: 13px;
        box-shadow: 0 14px 35px rgba(0,0,0,0.2);
    }
    .report .rh {
        color: #f1f1ef;
        font-weight: 800;
        font-size: 15px;
        margin: 18px 0 6px 0;
        padding-bottom: 4px;
        border-bottom: 1px solid #282b32;
    }
    .report .rh:first-child { margin-top: 0; }
    .report p { margin: 4px 0; font-size: 13px; }
    .report ul { margin: 4px 0 10px 18px; padding: 0; }
    .report li { color: #d7d7d4; font-size: 13px; line-height: 1.65; }
    .report b { color: #f1f1ef; }
    .report table.rt { border-collapse: collapse; width: 100%; margin: 8px 0; }
    .report table.rt th, .report table.rt td {
        border: 1px solid #282b32; padding: 6px 10px; font-size: 12px; text-align: left;
    }
    .report table.rt th { background: rgba(122, 34, 55, 0.25); color: #f1f1ef; }


    /* ======================================================
        SIDEBAR & TYPOGRAPHY
        ====================================================== */

    section[data-testid="stSidebar"] {
        background: linear-gradient(180deg, #0d0e11 0%, #090a0d 55%, #07080a 100%);
        border-right: 1px solid rgba(255, 255, 255, 0.065);
        box-shadow: 8px 0 35px rgba(0, 0, 0, 0.18);
    }

    section[data-testid="stSidebar"] > div {
        padding-top: 0.9rem;
    }

    h1 {
        color: #f5f4f1 !important;
        font-weight: 900 !important;
        letter-spacing: -1.8px !important;
    }

    h3 {
        color: #e5e3df !important;
        font-weight: 750 !important;
    }

    p, li {
        color: #c8c9cc;
        line-height: 1.75;
    }


    /* ======================================================
        INPUTS & BUTTONS
        ====================================================== */

    textarea, input {
        background: #0c0e12 !important;
        color: #f2f1ee !important;
        border: 1px solid #292c33 !important;
        border-radius: 10px !important;
        font-size: 13px !important;
    }

    div.stButton > button {
        min-height: 45px;
        border-radius: 10px;
        border: 1px solid rgba(177, 76, 101, 0.28);
        background: linear-gradient(105deg, #79283f 0%, #612536 48%, #49212d 100%);
        color: #ffffff;
        font-size: 12px;
        font-weight: 800;
        box-shadow: 0 10px 28px rgba(0, 0, 0, 0.22);
        transition: all 0.2s ease;
    }

    div.stButton > button:hover {
        background: linear-gradient(105deg, #893149 0%, #6c293b 50%, #542532 100%);
        transform: translateY(-1px);
    }

    div[data-testid="stDownloadButton"] button {
        min-height: 43px;
        border-radius: 10px;
        background: linear-gradient(145deg, #111318, #0d0f13);
        border: 1px solid #292c34;
        color: #d9dadc;
        font-size: 11px;
        font-weight: 750;
    }

    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# GENERIC HELPERS
# ============================================================

def get_api_key():
    """User-provided key (sidebar) first, then Streamlit secrets."""
    key = (st.session_state.get("user_groq_api_key") or "").strip()
    if key:
        return key
    try:
        return st.secrets.get("GROQ_API_KEY")
    except Exception:
        return None


def _retry_wait(msg):
    """Read 'try again in 6.2s' / '1m3.5s' / '350ms' from a Groq rate-limit message."""
    m = re.search(r"try again in\s*(?:(\d+)m)?\s*([\d.]+)s", msg)
    if m:
        return int(m.group(1) or 0) * 60 + float(m.group(2))
    m = re.search(r"try again in\s*([\d.]+)ms", msg)
    if m:
        return float(m.group(1)) / 1000.0
    return 6.0


def call_llm(api_key, system, user, max_tokens=900, temperature=0.2, retries=4):
    """One Groq call with automatic retry for rate limits / empty answers."""
    client = Groq(api_key=api_key)
    last_error = None
    for attempt in range(retries):
        try:
            response = client.chat.completions.create(
                model=MODEL,
                messages=[
                    {"role": "system", "content": system},
                    {"role": "user", "content": user},
                ],
                temperature=temperature,
                max_completion_tokens=max_tokens,
                reasoning_effort="low",
            )
            text = (response.choices[0].message.content or "").strip()
            if text:
                return text
            # reasoning tokens ate the budget -> give more room next time
            max_tokens = int(max_tokens * 1.5)
            last_error = RuntimeError("Model returned an empty response.")
        except Exception as e:
            last_error = e
            msg = str(e).lower()
            if "401" in msg or "invalid api key" in msg or "invalid_api_key" in msg:
                raise RuntimeError("AUTH: Invalid or missing Groq API key.")
            if "429" in msg or "rate limit" in msg:
                wait = _retry_wait(msg)
                if wait > 60:
                    raise RuntimeError(f"RATE: Groq limit reached, retry in about {int(wait)}s.")
                time.sleep(min(wait + 0.5, 45))
            else:
                time.sleep(2)
    raise last_error if last_error else RuntimeError("Unknown LLM error.")


def parse_json(text):
    """Extract a JSON object from model output (handles ```json fences)."""
    if not text:
        return None
    cleaned = re.sub(r"```(?:json)?", "", text).strip()
    start_i, end_i = cleaned.find("{"), cleaned.rfind("}")
    if start_i == -1 or end_i <= start_i:
        return None
    try:
        obj = json.loads(cleaned[start_i:end_i + 1])
        return obj if isinstance(obj, dict) else None
    except Exception:
        return None


def run_agent(log, api_key, key, role, user_prompt, max_tokens, as_json):
    """Run one agent (one LLM call) and record it in the execution log."""
    t0 = time.time()
    entry = {"key": key, "agent": role, "status": "done", "seconds": 0.0, "note": "", "fatal": False}

    if any(l.get("fatal") for l in log):
        entry.update(status="skipped", note="Skipped (earlier fatal error)")
        log.append(entry)
        return "", None

    system = (
        f"You are the {role} agent inside BusinessOps AI, a multi-agent business operations pipeline. "
        "Be specific to the user's actual request. No generic filler. Follow the output format exactly."
    )
    text, parsed = "", None
    try:
        text = call_llm(api_key, system, user_prompt, max_tokens=max_tokens)
        if as_json:
            parsed = parse_json(text)
            if parsed is None:
                entry.update(status="fallback", note="Invalid JSON, safe defaults used")
    except Exception as e:
        entry.update(status="failed", note=str(e)[:180], fatal=str(e).startswith(("AUTH", "RATE")))
        text, parsed = "", None

    entry["seconds"] = time.time() - t0
    log.append(entry)
    return text, parsed


def brief(data):
    """Shared context every agent receives (request + horizon + priority)."""
    tp, pr = data["time_period"], data["priority"]
    return (
        f"BUSINESS REQUEST:\n{data['request'][:2000]}\n\n"
        f"TIME HORIZON: {tp}\n-> {HORIZON_NOTES.get(tp, '')}\n"
        f"PRIORITY: {pr}\n-> {PRIORITY_NOTES.get(pr, '')}\n"
    )


def compact(obj, limit=1100):
    return json.dumps(obj, ensure_ascii=False)[:limit]


# ---------- normalizers (never trust raw LLM JSON) ----------

def _d(obj):
    return obj if isinstance(obj, dict) else {}


def _l(obj):
    return obj if isinstance(obj, list) else []


def _s(x, default=""):
    return str(x).strip() if x is not None and str(x).strip() else default


def _lvl(x, table, default):
    v = _s(x).lower()
    for name in table:
        if name in v:
            return name.capitalize()
    return default


def default_phases():
    return [
        {"name": "Setup & Alignment", "duration_pct": 20, "objective": "Align stakeholders and confirm scope.",
         "activities": ["Confirm scope and constraints", "Map resources and owners"], "owner": "Project Lead",
         "deliverable": "Approved plan"},
        {"name": "Core Execution", "duration_pct": 50, "objective": "Deliver the main operational changes.",
         "activities": ["Run core workflow", "Track dependencies and blockers"], "owner": "Operations Manager",
         "deliverable": "Working process"},
        {"name": "Optimization & QA", "duration_pct": 30, "objective": "Validate results and hand off.",
         "activities": ["Review KPIs", "Fix gaps and sign off"], "owner": "QA Lead",
         "deliverable": "Signed-off handoff"},
    ]


def norm_phases(obj):
    phases = []
    for p in _l(_d(obj).get("phases"))[:5]:
        if not isinstance(p, dict):
            continue
        try:
            pct = float(p.get("duration_pct", 0))
        except Exception:
            pct = 0
        phases.append({
            "name": _s(p.get("name"), "Phase"),
            "duration_pct": pct if pct > 0 else 20,
            "objective": _s(p.get("objective")),
            "activities": [_s(a) for a in _l(p.get("activities")) if _s(a)][:4],
            "owner": _s(p.get("owner"), "TBD"),
            "deliverable": _s(p.get("deliverable"), "-"),
        })
    return phases if len(phases) >= 2 else default_phases()


def norm_risks(obj):
    risks = []
    for r in _l(_d(obj).get("risks"))[:6]:
        if not isinstance(r, dict) or not _s(r.get("risk")):
            continue
        lk = _lvl(r.get("likelihood"), LIKELIHOOD, "Medium")
        im = _lvl(r.get("impact"), IMPACT, "Medium")
        score = LIKELIHOOD[lk.lower()] * IMPACT[im.lower()]
        risks.append({
            "risk": _s(r.get("risk")),
            "category": _s(r.get("category"), "Operational"),
            "likelihood": lk, "impact": im, "score": score,
            "mitigation": _s(r.get("mitigation"), "Define mitigation with owner."),
            "owner": _s(r.get("owner"), "TBD"),
        })
    if risks:
        return risks
    return [
        {"risk": "Resource bottlenecks", "category": "Operational", "likelihood": "Medium", "impact": "High",
         "score": 6, "mitigation": "Map resources early and keep a buffer.", "owner": "Operations Manager"},
        {"risk": "Timeline slippage", "category": "Timeline", "likelihood": "High", "impact": "High",
         "score": 9, "mitigation": "Weekly checkpoints and escalation path.", "owner": "Project Lead"},
        {"risk": "Communication gaps", "category": "Communication", "likelihood": "Medium", "impact": "Medium",
         "score": 4, "mitigation": "Single status channel and RACI.", "owner": "Project Lead"},
    ]


def norm_actions(obj):
    obj = _d(obj)
    out = {}
    for k in ("immediate", "next", "later"):
        out[k] = [_s(a) for a in _l(obj.get(k)) if _s(a)][:4]
    if not any(out.values()):
        out = {
            "immediate": ["Confirm scope and owners - Project Lead"],
            "next": ["Launch core workflow - Operations Manager"],
            "later": ["Review results and optimize - QA Lead"],
        }
    return out


def norm_kpis(obj):
    kpis = []
    for k in _l(_d(obj).get("kpis"))[:6]:
        if isinstance(k, dict) and _s(k.get("name")):
            kpis.append({
                "KPI": _s(k.get("name")), "Target": _s(k.get("target"), "-"),
                "How to measure": _s(k.get("measure"), "-"), "Frequency": _s(k.get("frequency"), "-"),
            })
    return kpis or [{"KPI": "On-time milestone completion", "Target": ">= 90%",
                     "How to measure": "Milestones done vs planned", "Frequency": "Weekly"}]


def norm_qa(obj):
    obj = _d(obj)
    if not obj:
        return {"score": None, "verdict": "QA not completed", "summary": "", "strengths": [], "gaps": [], "fixes": []}
    try:
        score = max(0, min(100, int(float(obj.get("score")))))
    except Exception:
        score = None
    return {
        "score": score,
        "verdict": _s(obj.get("verdict"), "Reviewed"),
        "summary": _s(obj.get("executive_summary")),
        "strengths": [_s(x) for x in _l(obj.get("strengths")) if _s(x)][:4],
        "gaps": [_s(x) for x in _l(obj.get("gaps")) if _s(x)][:4],
        "fixes": [_s(x) for x in _l(obj.get("fixes")) if _s(x)][:4],
    }


# ---------- schedule (real dates from the chosen horizon) ----------

def span_label(start_h, end_h, total_h):
    if total_h <= 48:
        return f"Hour {int(round(start_h))}-{int(round(end_h))}"
    if total_h >= 2160:
        w0 = int(start_h // 168) + 1
        return f"Week {w0}-{max(w0, int(round(end_h / 168)))}"
    d0 = int(start_h // 24) + 1
    return f"Day {d0}-{max(d0, int(round(end_h / 24)))}"


def build_schedule(phases, total_h, start_dt):
    total_pct = sum(p["duration_pct"] for p in phases) or 1
    cur, out = 0.0, []
    for p in phases:
        hours = total_h * p["duration_pct"] / total_pct
        out.append({
            **p,
            "start": start_dt + timedelta(hours=cur),
            "end": start_dt + timedelta(hours=cur + hours),
            "span": span_label(cur, cur + hours, total_h),
            "share": round(100 * p["duration_pct"] / total_pct),
        })
        cur += hours
    return out


def risk_level(risks):
    if not risks:
        return "N/A"
    avg = sum(r["score"] for r in risks) / len(risks)
    if avg < 3:
        return "Low"
    if avg < 5.5:
        return "Moderate"
    if avg < 8:
        return "Elevated"
    return "High"


# ---------- final markdown report ----------

def build_report_md(res):
    lines = []
    lines.append("## Executive Summary")
    lines.append(res["summary"] or "Executive summary unavailable (QA agent did not complete).")
    lines.append(f"- **Time horizon:** {res['time_period']}")
    lines.append(f"- **Priority:** {res['priority']}")
    lines.append(f"- **Generated:** {res['generated']}")

    lines.append("## Business Analysis")
    lines.append(res["analysis"])

    lines.append(f"## Recommended Workflow ({res['time_period']} Horizon)")
    for i, p in enumerate(res["schedule"], 1):
        lines.append(f"**Phase {i}: {p['name']}** ({p['span']} | {p['share']}% of timeline)")
        if p["objective"]:
            lines.append(f"- Objective: {p['objective']}")
        for a in p["activities"]:
            lines.append(f"- {a}")
        lines.append(f"- Owner: {p['owner']} | Deliverable: {p['deliverable']}")

    lines.append("## Risks & Mitigations")
    for r in res["risks"]:
        lines.append(
            f"- **{r['risk']}** ({r['category']}) - Likelihood: {r['likelihood']}, Impact: {r['impact']}. "
            f"Mitigation: {r['mitigation']} Owner: {r['owner']}"
        )

    lines.append(f"## Priority Actions (Priority: {res['priority']})")
    for label, k in (("Immediate", "immediate"), ("Next", "next"), ("Later", "later")):
        if res["actions"].get(k):
            lines.append(f"**{label}**")
            for a in res["actions"][k]:
                lines.append(f"- {a}")

    lines.append("## KPIs / Success Metrics")
    for k in res["kpis"]:
        lines.append(f"- **{k['KPI']}** - Target: {k['Target']} | Measure: {k['How to measure']} | Frequency: {k['Frequency']}")

    qa = res["qa"]
    lines.append("## QA Check")
    score = f"{qa['score']}/100" if qa["score"] is not None else "N/A"
    lines.append(f"- **Readiness score:** {score} - {qa['verdict']}")
    for label, key in (("Strength", "strengths"), ("Gap", "gaps"), ("Recommended fix", "fixes")):
        for item in qa[key]:
            lines.append(f"- {label}: {item}")
    return "\n".join(lines)


# ============================================================
# EXPORT HELPERS (PDF & DOCX)
# ============================================================

def _pdf_text(t):
    t = (t.replace("→", "->").replace("≥", ">=").replace("≤", "<=")
          .replace("‑", "-").replace("\u00a0", " "))
    t = t.encode("cp1252", "replace").decode("cp1252")
    t = html.escape(t, quote=False)
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)


def create_pdf(text_content):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    body = ParagraphStyle("ReportNormal", parent=styles["Normal"], fontSize=10, leading=14,
                          textColor=colors.HexColor("#222222"))
    h2 = ParagraphStyle("ReportH2", parent=styles["Heading2"], fontSize=13, leading=16, spaceBefore=10,
                        spaceAfter=4, textColor=colors.HexColor("#612536"))
    title = ParagraphStyle("ReportTitle", parent=styles["Title"], fontSize=18, textColor=colors.HexColor("#49212d"))
    bullet = ParagraphStyle("ReportBullet", parent=body, leftIndent=14, bulletIndent=3)

    story = [Paragraph("BusinessOps AI - Operations Report", title), Spacer(1, 8)]
    for raw in text_content.split("\n"):
        line = raw.strip()
        if not line:
            story.append(Spacer(1, 6))
        elif line.startswith("#"):
            story.append(Paragraph(_pdf_text(line.lstrip("#").strip()), h2))
        elif line.startswith(("- ", "* ")):
            story.append(Paragraph(_pdf_text(line[2:]), bullet, bulletText="•"))
        else:
            story.append(Paragraph(_pdf_text(line), body))
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def _docx_runs(paragraph, text):
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part)


def create_docx(text_content):
    doc = Document()
    doc.add_heading("BusinessOps AI - Operations Report", level=1)
    for raw in text_content.split("\n"):
        line = raw.strip()
        if not line:
            continue
        if line.startswith("#"):
            doc.add_heading(line.lstrip("#").strip(), level=2)
        elif line.startswith(("- ", "* ")):
            _docx_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            _docx_runs(doc.add_paragraph(), line)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# ---------- markdown -> safe HTML for the styled report box ----------

def md_to_html(md):
    out, in_list, table = [], False, []

    def inline(t):
        t = html.escape(t, quote=False)
        return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)

    def close_list():
        nonlocal in_list
        if in_list:
            out.append("</ul>")
            in_list = False

    def flush_table():
        nonlocal table
        rows = [[c.strip() for c in r.strip("|").split("|")]
                for r in table if not re.match(r"^\|?[\s:\-|]+\|?$", r)]
        if rows:
            out.append("<table class='rt'>")
            for i, r in enumerate(rows):
                tag = "th" if i == 0 else "td"
                out.append("<tr>" + "".join(f"<{tag}>{inline(c)}</{tag}>" for c in r) + "</tr>")
            out.append("</table>")
        table = []

    for raw in md.split("\n"):
        s = raw.strip()
        if s.startswith("|"):
            close_list()
            table.append(s)
            continue
        if table:
            flush_table()
        if not s:
            close_list()
        elif s.startswith("#"):
            close_list()
            out.append(f"<div class='rh'>{inline(s.lstrip('#').strip())}</div>")
        elif re.match(r"^([-*•]|\d+[.)])\s+", s):
            if not in_list:
                out.append("<ul>")
                in_list = True
            out.append(f"<li>{inline(re.sub(r'^([-*•]|\d+[.)])\s+', '', s))}</li>")
        else:
            close_list()
            out.append(f"<p>{inline(s)}</p>")
    close_list()
    if table:
        flush_table()
    return "".join(out)


# ============================================================
# CREWAI FLOW  (6 agents, each one real LLM call, chained)
# ============================================================

class BusinessOpsFlow(Flow):
    """
    cfg / ops_api_key / agent_log are plain attributes set by the UI before kickoff().
    Each agent receives the request, horizon, priority AND the output of previous agents.
    """

    @start()
    def intake(self):
        cfg = getattr(self, "cfg", {})
        return {
            "request": str(cfg.get("request", "")).strip(),
            "time_period": cfg.get("time_period", "30 Days (Short-term)"),
            "priority": cfg.get("priority", "Medium"),
            "outputs": {},
        }

    # ---- Agent 01: Business Analyst ----
    @listen(intake)
    def business_analysis(self, data):
        task = (
            brief(data)
            + "\nTASK: Analyze this business problem. Output exactly 5 markdown bullets with bold labels: "
              "**Objective**, **Stakeholders**, **Current Situation**, **Constraints**, **Expected Outcome**. "
              "Max 150 words total. Tailor to the horizon and priority. No headings, no tables."
        )
        text, _ = run_agent(self.agent_log, self.ops_api_key, "analysis", "Business Analyst", task, 700, False)
        analysis = text or "- Analysis unavailable: the Business Analyst agent did not return a result."
        return {**data, "outputs": {**data["outputs"], "analysis": analysis}}

    # ---- Agent 02: Operations Planner ----
    @listen(business_analysis)
    def operations_planning(self, data):
        task = (
            brief(data)
            + f"\nBUSINESS ANALYSIS:\n{data['outputs']['analysis'][:1200]}\n\n"
              "TASK: Design the execution workflow as 3 to 5 sequential phases that fit the time horizon and "
              "priority. Return ONLY valid JSON, no extra text, in this shape:\n"
              '{"phases":[{"name":"short name","duration_pct":20,"objective":"one sentence",'
              '"activities":["step","step","step"],"owner":"role","deliverable":"what is produced"}]}\n'
              "duration_pct values must add up to 100. Use roles, not people names."
        )
        _, parsed = run_agent(self.agent_log, self.ops_api_key, "operations", "Operations Planner", task, 1200, True)
        return {**data, "outputs": {**data["outputs"], "phases": norm_phases(parsed)}}

    # ---- Agent 03: Risk Manager ----
    @listen(operations_planning)
    def risk_management(self, data):
        task = (
            brief(data)
            + f"\nPLANNED PHASES:\n{compact(data['outputs']['phases'])}\n\n"
              "TASK: Identify 4 to 6 concrete risks for THIS plan (operational, people, technology, communication, "
              "timeline). Return ONLY valid JSON:\n"
              '{"risks":[{"risk":"short title","category":"Operational","likelihood":"Low|Medium|High",'
              '"impact":"Low|Medium|High|Critical","mitigation":"specific action","owner":"role"}]}'
        )
        _, parsed = run_agent(self.agent_log, self.ops_api_key, "risk", "Risk Manager", task, 1100, True)
        return {**data, "outputs": {**data["outputs"], "risks": norm_risks(parsed)}}

    # ---- Agent 04: Action Planner ----
    @listen(risk_management)
    def action_planning(self, data):
        top_risks = sorted(data["outputs"]["risks"], key=lambda r: -r["score"])[:3]
        task = (
            brief(data)
            + f"\nPHASES:\n{compact(data['outputs']['phases'], 900)}\n"
              f"TOP RISKS:\n{compact(top_risks, 700)}\n\n"
              "TASK: Create prioritized next actions. Each action must be one line formatted as "
              "'Action - Owner role - Deadline'. Deadlines must respect the time horizon and priority "
              "(hours for urgent/immediate, days or weeks otherwise). Return ONLY valid JSON:\n"
              '{"immediate":["..",".."],"next":["..",".."],"later":["..",".."]}'
        )
        _, parsed = run_agent(self.agent_log, self.ops_api_key, "actions", "Action Planner", task, 900, True)
        return {**data, "outputs": {**data["outputs"], "actions": norm_actions(parsed)}}

    # ---- Agent 05: KPI Designer ----
    @listen(action_planning)
    def kpi_design(self, data):
        task = (
            brief(data)
            + f"\nBUSINESS ANALYSIS:\n{data['outputs']['analysis'][:700]}\n"
              f"PHASES:\n{compact(data['outputs']['phases'], 800)}\n\n"
              "TASK: Define 4 to 6 measurable KPIs with realistic targets for this horizon. "
              "Return ONLY valid JSON:\n"
              '{"kpis":[{"name":"KPI name","target":"numeric target","measure":"how it is measured",'
              '"frequency":"how often"}]}'
        )
        _, parsed = run_agent(self.agent_log, self.ops_api_key, "kpi", "KPI Designer", task, 900, True)
        return {**data, "outputs": {**data["outputs"], "kpis": norm_kpis(parsed)}}

    # ---- Agent 06: QA Auditor ----
    @listen(kpi_design)
    def quality_control(self, data):
        o = data["outputs"]
        package = {
            "analysis": o["analysis"][:700],
            "phases": [{"name": p["name"], "pct": p["duration_pct"], "activities": p["activities"]} for p in o["phases"]],
            "risks": [{"risk": r["risk"], "L": r["likelihood"], "I": r["impact"]} for r in o["risks"]],
            "actions": o["actions"],
            "kpis": [{"name": k["KPI"], "target": k["Target"]} for k in o["kpis"]],
        }
        task = (
            brief(data)
            + f"\nFULL PLAN TO AUDIT:\n{compact(package, 3000)}\n\n"
              "TASK: Audit the plan for practicality, completeness, consistency with the horizon/priority and "
              "measurability. Be honest; do not give a perfect score by default. Return ONLY valid JSON:\n"
              '{"score":0-100,"verdict":"Ready|Ready with fixes|Needs revision",'
              '"executive_summary":"60-90 word summary of the whole plan for executives",'
              '"strengths":["..",".."],"gaps":["..",".."],"fixes":["..",".."]}'
        )
        _, parsed = run_agent(self.agent_log, self.ops_api_key, "qa", "QA Auditor", task, 1000, True)
        return {**data, "outputs": {**o, "qa": norm_qa(parsed)}}

    # ---- Final assembly (no LLM call) ----
    @listen(quality_control)
    def final_report(self, data):
        if not data["request"]:
            return {"error": "Please enter a business request."}

        o = data["outputs"]
        total_h = HORIZON_HOURS.get(data["time_period"], 720)
        res = {
            "request": data["request"],
            "time_period": data["time_period"],
            "priority": data["priority"],
            "generated": datetime.now().strftime("%d %b %Y, %H:%M"),
            "total_h": total_h,
            "analysis": o["analysis"],
            "schedule": build_schedule(o["phases"], total_h, datetime.now()),
            "risks": o["risks"],
            "actions": o["actions"],
            "kpis": o["kpis"],
            "qa": o["qa"],
            "summary": o["qa"]["summary"],
            "log": list(self.agent_log),
        }
        res["report"] = build_report_md(res)
        return res


# ============================================================
# SIDEBAR
# ============================================================

with st.sidebar:
    st.markdown("## ⚡ BusinessOps AI")
    st.markdown("<p style='font-size:12px; color:#969aa3;'>Autonomous Business Process Intelligence Platform</p>", unsafe_allow_html=True)
    st.divider()

    st.markdown("### Technology")
    st.write("🐍 Python 3.12")
    st.write("🎨 Streamlit")
    st.write("🤖 CrewAI Flow")
    st.write("⚡ Groq API")
    st.write("🧠 GPT-OSS 20B")
    st.divider()

    st.markdown("### API Key")
    st.text_input(
        "Groq API key (optional)",
        type="password",
        key="user_groq_api_key",
        help="Leave empty to use GROQ_API_KEY from Streamlit secrets.",
    )
    if st.button("🗑 Clear results", **FULL):
        for k in ("result", "chat_history", "flash"):
            st.session_state.pop(k, None)
        st.rerun()
    st.divider()
    st.caption("Free-tier friendly architecture")


# ============================================================
# HERO SECTION
# ============================================================

st.markdown(
    """
    <div class="hero">
        <span class="badge">AUTONOMOUS BUSINESS INTELLIGENCE</span>
        <h1>BusinessOps AI</h1>
        <p>Transform complex business requests into structured operational plans, risk controls, priority actions, and measurable outcomes using an agentic AI workflow.</p>
    </div>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# WORKFLOW STAGES (live status after a run)
# ============================================================

st.markdown("### Agentic Workflow")

stages = [
    ("01", "Business Analyst", "Understands the business problem, objective, stakeholders and constraints.", "analysis"),
    ("02", "Operations Planner", "Converts the problem into an executable operational workflow.", "operations"),
    ("03", "Risk Manager", "Identifies implementation risks and practical mitigation strategies.", "risk"),
    ("04", "Action Planner", "Converts recommendations into prioritized next actions.", "actions"),
    ("05", "KPI Designer", "Defines measurable outcomes and success indicators.", "kpi"),
    ("06", "QA Auditor", "Performs a final quality and consistency review.", "qa"),
]


def status_html(entry):
    if not entry:
        return "<div class='stage-status idle'>● Standby</div>"
    t = entry["seconds"]
    s = entry["status"]
    if s == "done":
        return f"<div class='stage-status ok'>✔ Completed · {t:.1f}s</div>"
    if s == "fallback":
        return f"<div class='stage-status warn'>⚠ Defaults used · {t:.1f}s</div>"
    if s == "skipped":
        return "<div class='stage-status warn'>— Skipped</div>"
    return "<div class='stage-status fail'>✖ Failed</div>"


current = st.session_state.get("result") or {}
log_map = {l["key"]: l for l in current.get("log", [])}

cols = st.columns(3)
for i, (number, title, description, key) in enumerate(stages):
    with cols[i % 3]:
        st.markdown(
            f"<div class='stage'><div class='stage-number'>{number} / 06</div>"
            f"<div class='stage-title'>{title}</div>"
            f"<div class='stage-text'>{description}</div>"
            f"{status_html(log_map.get(key))}</div>",
            unsafe_allow_html=True,
        )


# ============================================================
# METRICS SECTION
# ============================================================

st.write("")
m1, m2, m3, m4 = st.columns(4)
with m1:
    st.metric(label="WORKFLOW STAGES", value="06")
with m2:
    st.metric(label="LLM CALLS / RUN", value="06")
with m3:
    st.metric(label="MODEL", value="20B")
with m4:
    st.metric(label="DEPLOYMENT", value="CLOUD")
st.write("")


# ============================================================
# BUSINESS REQUEST INPUT & ADVANCED CONTROLS
# ============================================================

st.markdown("### Business Request & Configuration")

col_c1, col_c2 = st.columns(2)
with col_c1:
    time_period = st.selectbox(
        "⏱️ Execution Time Horizon",
        ["Immediate (24-48 Hours)", "30 Days (Short-term)", "90 Days (Quarterly)", "6 Months (Strategic)"],
    )
with col_c2:
    priority = st.selectbox(
        "🔥 Priority Level",
        ["Critical / Urgent", "High", "Medium", "Low"],
    )

sample = st.selectbox(
    "Quick scenario",
    [
        "Custom request",
        "Employee Onboarding",
        "Software Rollout",
        "Office Relocation",
        "Customer Support Improvement",
    ],
)

SAMPLES = {
    "Employee Onboarding": "Our company is growing quickly and new employees are having difficulty completing HR, IT, security and department onboarding. Design a better onboarding process.",
    "Software Rollout": "A company is introducing a new internal software platform. Employees need training, communication, migration support and a controlled rollout plan.",
    "Office Relocation": "Our organization is moving to a new office. We need a plan covering employees, IT infrastructure, vendors, communication, facilities and business continuity.",
    "Customer Support Improvement": "Customer support response times are increasing and customers are complaining about inconsistent answers. Create an improved support operations workflow.",
}

request = st.text_area(
    "Describe your business problem or process",
    value=SAMPLES.get(sample, ""),
    height=170,
    placeholder="Example: Our company wants to improve employee onboarding...",
    key=f"request_{sample}",
)


# ============================================================
# RUN BUTTON  (runs the 6 agents, stores result, re-renders page)
# ============================================================

if st.button("⚡ RUN BUSINESSOPS AI", type="primary", **FULL):
    api_key = get_api_key()
    if not request.strip():
        st.warning("Please enter a business request first.")
    elif not api_key:
        st.error("GROQ_API_KEY is missing. Enter it in the sidebar or add it in Streamlit Cloud → Settings → Secrets.")
    else:
        flow = BusinessOpsFlow()
        flow.cfg = {"request": request, "time_period": time_period, "priority": priority}
        flow.ops_api_key = api_key
        flow.agent_log = []

        with st.spinner("6 agents are working: Analyst → Planner → Risk → Actions → KPIs → QA (30-90 seconds on free tier)..."):
            try:
                result = flow.kickoff()
            except Exception as e:
                result = {"error": f"Flow failed: {e}"}

        if not isinstance(result, dict):
            result = {"error": "Flow returned no result. Please try again."}

        if "error" in result:
            st.error(result["error"])
        else:
            ok = sum(1 for l in result["log"] if l["status"] in ("done", "fallback"))
            if ok == 0:
                notes = " | ".join(f"{l['agent']}: {l['note']}" for l in result["log"] if l["note"])
                st.error(f"All agents failed. {notes}")
            else:
                st.session_state["result"] = result
                st.session_state["chat_history"] = []
                st.session_state["flash"] = (
                    f"Business workflow completed: {ok}/6 agents finished in "
                    f"{sum(l['seconds'] for l in result['log']):.0f}s."
                )
                st.rerun()


# ============================================================
# INTERACTIVE DASHBOARD  (rendered from session state, so it
# survives chat messages and download clicks)
# ============================================================

def render_dashboard(res):
    flash = st.session_state.pop("flash", None)
    if flash:
        st.success(flash)

    st.write("")
    st.markdown("---")
    st.markdown("### 🎛️ Interactive Intelligence Dashboard")

    tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs([
        "📊 Operational Overview",
        "🗺 Milestone Roadmap",
        "🛡️ Risk Operations",
        "✅ Actions & KPIs",
        "📝 Full Intelligence Report",
        "🤖 Agent Execution Log",
    ])

    qa = res["qa"]
    sched = res["schedule"]

    # ---------- Overview ----------
    with tab1:
        om1, om2, om3, om4 = st.columns(4)
        with om1:
            st.metric(label="TIME HORIZON", value=HORIZON_SHORT.get(res["time_period"], res["time_period"]))
        with om2:
            st.metric(label="PRIORITY RATING", value=res["priority"].split(" ")[0])
        with om3:
            st.metric(label="RISK LEVEL", value=risk_level(res["risks"]))
        with om4:
            st.metric(label="EXECUTION READINESS", value=f"{qa['score']}%" if qa["score"] is not None else "N/A")

        on1, on2, on3, on4 = st.columns(4)
        with on1:
            st.metric(label="PHASES", value=len(sched))
        with on2:
            st.metric(label="RISKS IDENTIFIED", value=len(res["risks"]))
        with on3:
            st.metric(label="PRIORITY ACTIONS", value=sum(len(v) for v in res["actions"].values()))
        with on4:
            st.metric(label="KPIs DEFINED", value=len(res["kpis"]))

        st.write("")
        st.markdown("#### Executive Summary")
        st.write(res["summary"] or "Executive summary unavailable (QA agent did not complete).")
        st.markdown("#### Business Analysis")
        st.markdown(res["analysis"])

        if qa["gaps"] or qa["fixes"]:
            qc1, qc2 = st.columns(2)
            with qc1:
                if qa["gaps"]:
                    st.warning("**QA Gaps**\n\n" + "\n\n".join(f"• {g}" for g in qa["gaps"]))
            with qc2:
                if qa["fixes"]:
                    st.info("**Recommended Fixes**\n\n" + "\n\n".join(f"• {f}" for f in qa["fixes"]))

    # ---------- Roadmap ----------
    with tab2:
        st.markdown("#### 📅 Interactive Milestone Roadmap (Gantt View)")
        gantt = pd.DataFrame([
            {"Task": f"{i}. {p['name']}", "Start": p["start"], "Finish": p["end"], "Window": p["span"]}
            for i, p in enumerate(sched, 1)
        ])
        fig = px.timeline(
            gantt, x_start="Start", x_end="Finish", y="Task", color="Task", hover_data=["Window"],
            color_discrete_sequence=["#79283f", "#a13a55", "#b14c65", "#6a3a4a", "#c97a8b"],
        )
        fig.update_yaxes(autorange="reversed", title=None)
        fig.update_xaxes(title=None)
        fig.update_layout(
            paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)", font_color="#eeeeec",
            margin=dict(t=10, b=10, l=10, r=10), height=110 + 60 * len(sched), showlegend=False,
        )
        st.plotly_chart(fig, **FULL)

        painters = [st.info, st.warning, st.success]
        per_row = 3
        for row_start in range(0, len(sched), per_row):
            row = sched[row_start:row_start + per_row]
            r_cols = st.columns(per_row)
            for j, p in enumerate(row):
                idx = row_start + j
                acts = "\n\n".join(f"• {a}" for a in p["activities"])
                with r_cols[j]:
                    painters[idx % 3](
                        f"**Phase {idx + 1}: {p['name']}**\n\n*{p['span']} · {p['share']}% of timeline*\n\n"
                        f"{p['objective']}\n\n{acts}\n\n**Owner:** {p['owner']}\n\n**Deliverable:** {p['deliverable']}"
                    )

    # ---------- Risks ----------
    with tab3:
        risks = sorted(res["risks"], key=lambda r: -r["score"])
        rc = st.columns(3)
        with rc[0]:
            st.error("**Top Risk**\n\n" + f"{risks[0]['risk']}\n\nLikelihood: {risks[0]['likelihood']} · Impact: {risks[0]['impact']}")
        with rc[1]:
            st.warning("**Primary Mitigation**\n\n" + risks[0]["mitigation"])
        with rc[2]:
            st.success("**Risk Owner**\n\n" + risks[0]["owner"])

        st.write("")
        st.markdown("#### 🎯 Risk Impact vs Likelihood Matrix")
        rdf = pd.DataFrame([{
            "Risk Factor": r["risk"], "Category": r["category"], "Likelihood": r["likelihood"],
            "Impact": r["impact"], "Score": r["score"],
            "Status": "Active - escalate" if r["score"] >= 8 else ("Monitor" if r["score"] >= 4 else "Controlled"),
            "Mitigation": r["mitigation"], "Owner": r["owner"],
        } for r in risks])
        st.dataframe(rdf, hide_index=True, **FULL)

        bar = px.bar(rdf, x="Score", y="Risk Factor", orientation="h", color="Score",
                     color_continuous_scale=["#49212d", "#79283f", "#c97a8b"], range_x=[0, 12])
        bar.update_yaxes(autorange="reversed", title=None)
        bar.update_layout(
            paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)", font_color="#eeeeec",
            margin=dict(t=10, b=10, l=10, r=10), height=80 + 50 * len(rdf), coloraxis_showscale=False,
        )
        st.plotly_chart(bar, **FULL)

    # ---------- Actions & KPIs ----------
    with tab4:
        st.markdown(f"#### 🚀 Priority Actions ({res['priority']})")
        a1, a2, a3 = st.columns(3)
        for col, painter, label, key in (
            (a1, st.error, "Immediate", "immediate"),
            (a2, st.warning, "Next", "next"),
            (a3, st.success, "Later", "later"),
        ):
            with col:
                items = res["actions"].get(key) or ["-"]
                painter(f"**{label}**\n\n" + "\n\n".join(f"• {x}" for x in items))
        st.write("")
        st.markdown("#### 📈 KPIs / Success Metrics")
        st.dataframe(pd.DataFrame(res["kpis"]), hide_index=True, **FULL)

    # ---------- Full report ----------
    with tab5:
        st.markdown(f"<div class='report'>{md_to_html(res['report'])}</div>", unsafe_allow_html=True)
        st.write("")
        dl1, dl2, dl3 = st.columns(3)
        with dl1:
            st.download_button(
                "📥 Download Text (.txt)", data=res["report"],
                file_name="businessops_report.txt", mime="text/plain", **FULL,
            )
        with dl2:
            st.download_button(
                "📥 Download PDF (.pdf)", data=create_pdf(res["report"]),
                file_name="businessops_report.pdf", mime="application/pdf", **FULL,
            )
        with dl3:
            st.download_button(
                "📥 Download Word (.docx)", data=create_docx(res["report"]),
                file_name="businessops_report.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", **FULL,
            )

    # ---------- Agent log ----------
    with tab6:
        names = {k: t for _, t, _, k in stages}
        label = {"done": "✔ Completed", "fallback": "⚠ Defaults used", "failed": "✖ Failed", "skipped": "— Skipped"}
        ldf = pd.DataFrame([{
            "Agent": names.get(l["key"], l["agent"]), "Status": label.get(l["status"], l["status"]),
            "Time (s)": round(l["seconds"], 1), "Note": l["note"] or "-",
        } for l in res["log"]])
        st.dataframe(ldf, hide_index=True, **FULL)
        st.caption(f"Total pipeline time: {sum(l['seconds'] for l in res['log']):.1f}s · Generated {res['generated']}")


if st.session_state.get("result"):
    render_dashboard(st.session_state["result"])


# ============================================================
# CHAT WITH YOUR OPS PLAN
# ============================================================

if st.session_state.get("result"):
    st.write("")
    st.markdown("---")
    st.markdown("### 💬 Chat with your Ops Plan")
    st.markdown("<p style='font-size:12px; color:#969aa3;'>Ask follow-up questions, request specific sprint breakdowns, or query bottlenecks from your generated report.</p>", unsafe_allow_html=True)

    if "chat_history" not in st.session_state:
        st.session_state["chat_history"] = []

    for message in st.session_state["chat_history"]:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

    if user_query := st.chat_input("Ask a question about your operational plan..."):
        st.session_state["chat_history"].append({"role": "user", "content": user_query})
        with st.chat_message("user"):
            st.markdown(user_query)

        with st.chat_message("assistant"):
            with st.spinner("Analyzing operational plan..."):
                api_key = get_api_key()
                if not api_key:
                    chat_response = "API key missing. Enter it in the sidebar or configure GROQ_API_KEY in Streamlit secrets."
                else:
                    try:
                        history = "\n".join(
                            f"{m['role'].upper()}: {m['content']}" for m in st.session_state["chat_history"][-7:-1]
                        )
                        chat_prompt = (
                            "Answer the follow-up question using ONLY the operations report below. "
                            "If the report does not cover it, say so and suggest a practical next step.\n\n"
                            f"REPORT:\n{st.session_state['result']['report']}\n\n"
                            f"RECENT CHAT:\n{history or '(none)'}\n\n"
                            f"QUESTION:\n{user_query}\n\n"
                            "Reply concisely, practically and professionally."
                        )
                        chat_response = call_llm(
                            api_key,
                            "You are a professional business operations assistant answering queries about a generated report.",
                            chat_prompt, max_tokens=700, temperature=0.3,
                        )
                    except Exception as e:
                        chat_response = f"Could not generate chat response. Error: {e}"

            st.markdown(chat_response)
            st.session_state["chat_history"].append({"role": "assistant", "content": chat_response})


# ============================================================
# FOOTER
# ============================================================

st.divider()
st.markdown(
    """
    <div style="text-align:center; color:#667386; font-size:12px;">
        BusinessOps AI • Autonomous Business Process Intelligence • CrewAI + Groq • Free-tier deployment architecture
    </div>
    """,
    unsafe_allow_html=True,
)
